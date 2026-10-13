import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette: Pure White, Deep Charcoal, Professional Slate Blue
    BG_WHITE = RGBColor(255, 255, 255)
    TEXT_TITLE = RGBColor(15, 23, 42)      # Deep Charcoal
    TEXT_ACCENT = RGBColor(3, 105, 161)    # Deep Professional Blue
    TEXT_BODY = RGBColor(51, 65, 85)       # High-Contrast Slate
    TEXT_MUTED = RGBColor(100, 116, 139)   # Medium Gray
    CARD_BG = RGBColor(248, 250, 252)      # Soft Clean Off-White
    CARD_BORDER = RGBColor(203, 213, 225)  # Light Subtle Gray Border
    TAG_BG = RGBColor(224, 242, 254)       # Light Blue Tint
    TAG_TEXT = RGBColor(3, 105, 161)       # Blue Tag Text

    def set_slide_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_WHITE

    def add_header(slide, step_number, title_text, subtitle_text):
        # Header Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        if step_number:
            run_step = p1.add_run()
            run_step.text = step_number + "  "
            run_step.font.name = "Arial"
            run_step.font.size = Pt(14)
            run_step.font.bold = True
            run_step.font.color.rgb = TEXT_ACCENT

        run_title = p1.add_run()
        run_title.text = title_text
        run_title.font.name = "Arial"
        run_title.font.size = Pt(28)
        run_title.font.bold = True
        run_title.font.color.rgb = TEXT_TITLE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        run_sub = p2.add_run()
        run_sub.text = subtitle_text
        run_sub.font.name = "Arial"
        run_sub.font.size = Pt(15)
        run_sub.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: JUDUL UTAMA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Big Title Box in Center
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    r = p.add_run()
    r.text = "ThermalSight v1.8.0"
    r.font.name = "Arial"
    r.font.size = Pt(44)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p2 = tf1.add_paragraph()
    p2.space_before = Pt(12)
    r2 = p2.add_run()
    r2.text = "Alur Kerja Pengolahan & Analisis Foto Panas (Termal) Telapak Kaki"
    r2.font.name = "Arial"
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = TEXT_ACCENT

    p3 = tf1.add_paragraph()
    p3.space_before = Pt(16)
    r3 = p3.add_run()
    r3.text = "Panduan Sederhana & Lengkap: Dari Foto Kamera FLIR hingga Menjadi Hasil Analisis Siap Pakai"
    r3.font.name = "Arial"
    r3.font.size = Pt(16)
    r3.font.color.rgb = TEXT_BODY

    p4 = tf1.add_paragraph()
    p4.space_before = Pt(36)
    r4 = p4.add_run()
    r4.text = "Disusun secara ringkas, jelas, dan mudah dipahami untuk pembaca umum maupun klinisi."
    r4.font.name = "Arial"
    r4.font.size = Pt(13)
    r4.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: INPUT SEPERTI APA?
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "LANGKAH 1", "Data Masukan (Input Gambar)", "Jenis foto yang dimasukkan ke dalam aplikasi ThermalSight")

    # 3 Column Cards
    col_w = Inches(3.64)
    col_gap = Inches(0.4)
    y_top = Inches(2.2)
    card_h = Inches(4.5)

    c1_left = Inches(0.8)
    c2_left = c1_left + col_w + col_gap
    c3_left = c2_left + col_w + col_gap

    # Card 1: Foto Kamera FLIR
    box1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, y_top, col_w, card_h)
    box1.fill.solid()
    box1.fill.fore_color.rgb = CARD_BG
    box1.line.color.rgb = CARD_BORDER
    box1.line.width = Pt(1)
    tf = box1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "1. Foto Kamera Panas (FLIR)\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Menggunakan foto asli dari kamera termal FLIR (format .jpg, .png, atau .tiff).\n\n• Bukan sekadar foto warna biasa, melainkan foto khusus yang menyimpan data derajat panas (termal)."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # Card 2: Foto Tunggal atau Folder
    box2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, y_top, col_w, card_h)
    box2.fill.solid()
    box2.fill.fore_color.rgb = CARD_BG
    box2.line.color.rgb = CARD_BORDER
    box2.line.width = Pt(1)
    tf = box2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "2. Bebas Pilih Mode Masukan\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Bisa memasukkan 1 foto tunggal untuk diperiksa cepat.\n\n• Atau memasukkan 1 folder berisi puluhan foto urutan pasien (mode sekuens).\n\n• Bisa juga langsung geser dan lepas (drag & drop) foto ke layar."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # Card 3: Otomatis Menyaring Folder
    box3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_left, y_top, col_w, card_h)
    box3.fill.solid()
    box3.fill.fore_color.rgb = CARD_BG
    box3.line.color.rgb = CARD_BORDER
    box3.line.width = Pt(1)
    tf = box3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "3. Saringan Otomatis Pintar\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Jika memilih folder yang di dalamnya ada subfolder hasil sebelumnya (_Result atau _analysis), aplikasi otomatis mengabaikannya.\n\n• Hanya foto utama di folder terluar yang dimuat, sehingga tidak akan tercampur aduk."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 3: DIPROSES FLIR SEPERTI APA?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "LANGKAH 2", "Pembersihan Gambar (Proses FLIR Sensor Asli)", "Cara aplikasi membersihkan gambar agar foto kaki menjadi bersih alami")

    # 2 Wide Cards Side by Side
    w_card = Inches(5.66)
    c1 = Inches(0.8)
    c2 = c1 + w_card + Inches(0.4)

    # Left: Cara Lama vs Masalah
    box_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1, y_top, w_card, card_h)
    box_left.fill.solid()
    box_left.fill.fore_color.rgb = CARD_BG
    box_left.line.color.rgb = CARD_BORDER
    box_left.line.width = Pt(1)
    tf = box_left.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.35)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Tantangan Foto Bawaan Kamera FLIR:\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(12)
    r = p_body.add_run()
    r.text = "Kamera FLIR biasanya menaruh banyak aksesoris yang menutupi gambar:\n\n" \
             "• Tulisan suhu di pojok (contoh: teks '35.0 °C')\n" \
             "• Logo tulisan merek 'FLIR' di pojok bawah\n" \
             "• Mistar skala warna (colorbar) di tepi kanan\n" \
             "• Tanda silang (titik fokus) di tengah gambar\n\n" \
             "Jika dihapus dengan cara menambal biasa (blur/inpaint), gambar telapak kaki akan menjadi buram dan tidak akurat."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # Right: Solusi v1.8.0 Mode Mula-Mula
    box_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2, y_top, w_card, card_h)
    box_right.fill.solid()
    box_right.fill.fore_color.rgb = CARD_BG
    box_right.line.color.rgb = TEXT_ACCENT
    box_right.line.width = Pt(2)
    tf = box_right.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.35)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Solusi v1.8.0: Mode Mula-mula (Sensor Murni)\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_ACCENT

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(12)
    r = p_body.add_run()
    r.text = "Aplikasi membaca angka suhu mentah langsung dari dalam file (menggunakan pustaka 'flyr' dan rumus fisika Planck):\n\n" \
             "✓ Tidak ada tambalan blur atau coretan buatan sama sekali.\n" \
             "✓ Logo, tulisan '35°C', dan mistar warna hilang secara alami.\n" \
             "✓ Nilai suhu setiap titik dipertahankan 100% akurat sesuai pembacaan sensor kamera.\n\n" \
             "Hasil: Anda memperoleh gambar telapak kaki yang bersih total, bening, dan siap dianalisis."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 4: ALUR APLIKASI SEPERTI APA?
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "LANGKAH 3", "Alur Kerja Pemakaian Aplikasi (4 Tahap Praktis)", "Langkah demi langkah yang dilakukan pengguna dari awal sampai selesai")

    # 4 Steps Horizontal / Grid
    step_w = Inches(2.7)
    step_gap = Inches(0.3)
    step_h = Inches(4.5)

    steps_data = [
        ("TAHAP 1", "Unggah Foto", "Pilih 1 foto tunggal atau 1 folder sekuens pasien.\n\nLayar akan menampilkan proses pembacaan file dengan aman."),
        ("TAHAP 2", "Kalibrasi Ukuran", "Tarik kotak di atas kartu (misal kartu 10×10 cm) atau tarik garis penggaris.\n\nTujuannya: Agar ukuran di layar sama persis dengan ukuran kaki nyata (px ke cm)."),
        ("TAHAP 3", "Tandai Titik Kaki", "Cukup klik 1 kali pada ibu jari atau telapak kaki.\n\nLingkaran otomatis terbentuk dan arah rambat panas langsung dihitung."),
        ("TAHAP 4", "Simpan & Ekspor", "Klik tombol 'Save & Export'.\n\nAplikasi langsung membuatkan gambar hitam-putih, tabel angka excel, dan file arsip lengkap.")
    ]

    for i, (tag, stitle, sdesc) in enumerate(steps_data):
        left_pos = Inches(0.8) + i * (step_w + step_gap)
        s_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, y_top, step_w, step_h)
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = CARD_BG
        s_box.line.color.rgb = CARD_BORDER
        s_box.line.width = Pt(1)

        tf = s_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)

        p = tf.paragraphs[0]
        r_tag = p.add_run()
        r_tag.text = tag + "\n"
        r_tag.font.name = "Arial"
        r_tag.font.size = Pt(12)
        r_tag.font.bold = True
        r_tag.font.color.rgb = TEXT_ACCENT

        r_title = p.add_run()
        r_title.text = stitle + "\n"
        r_title.font.name = "Arial"
        r_title.font.size = Pt(18)
        r_title.font.bold = True
        r_title.font.color.rgb = TEXT_TITLE

        p_body = tf.add_paragraph()
        p_body.space_before = Pt(8)
        r_desc = p_body.add_run()
        r_desc.text = sdesc
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(14)
        r_desc.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 5: MENGHASILKAN 9x9 GRID SEPERTI APA?
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "LANGKAH 4", "Standarisasi Lingkaran Grid 9×9", "Bagaimana lingkaran pengamatan dibuat dan mengapa harus 9×9 sel kotak")

    # 3 Column Cards
    box1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, y_top, col_w, card_h)
    box1.fill.solid()
    box1.fill.fore_color.rgb = CARD_BG
    box1.line.color.rgb = CARD_BORDER
    box1.line.width = Pt(1)
    tf = box1.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "1. Ukuran Pas 9×9 Kotak\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Di dalam analisis ilmiah, telapak kaki dibagi menjadi kotak-kotak kecil teratur.\n\n• Lingkaran penanda titik kaki distandarkan berdiameter tepat 9 kotak (radius 4.5 kotak).\n\n• Tujuannya: Agar perhitungan suhu rata-rata di area tersebut selalu konsisten dan adil di semua pasien."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    box2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, y_top, col_w, card_h)
    box2.fill.solid()
    box2.fill.fore_color.rgb = CARD_BG
    box2.line.color.rgb = CARD_BORDER
    box2.line.width = Pt(1)
    tf = box2.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "2. Menyesuaikan Jarak Kamera\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_TITLE

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Saat memotret, jarak kamera ke kaki bisa berbeda-beda (ada yang agak dekat, ada yang agak jauh).\n\n• Berkat kalibrasi ukuran di awal, aplikasi otomatis mengatur besarnya lingkaran di layar.\n\n• Ukuran aslinya di kulit pasien tetap sama persis: 1.2 sentimeter."
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    box3 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_left, y_top, col_w, card_h)
    box3.fill.solid()
    box3.fill.fore_color.rgb = CARD_BG
    box3.line.color.rgb = TEXT_ACCENT
    box3.line.width = Pt(2)
    tf = box3.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.3)

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "3. Otomatis Naik Kelas\n"
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = TEXT_ACCENT

    p_body = tf.add_paragraph()
    p_body.space_before = Pt(10)
    r = p_body.add_run()
    r.text = "• Jika Anda membuka file catatan dari aplikasi versi lama (v1.6 / v1.7) yang lingkarannya masih kecil (6×6):\n\n✓ Aplikasi otomatis memperbesar lingkaran ke standar baru 9×9.\n✓ Posisi titik tengah kaki tidak akan bergeser sedikit pun.\n✓ Tidak perlu repot menandai ulang dari awal!"
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = TEXT_BODY

    # ==========================================
    # SLIDE 6: HASIL AKHIR APA SAJA?
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "LANGKAH 5", "Hasil Akhir yang Didapatkan (Output Analisis)", "Dokumen dan gambar lengkap yang langsung siap digunakan untuk laporan")

    # 4 Output Cards
    out_data = [
        ("Gambar Telapak Kaki Jurnal (Fig. 1)", 
         "• Latar belakang putih-hitam kontras tinggi.\n• Lingkaran warna suhu tajam di area yang diperiksa.\n• Pilihan garis halus standar jurnal internasional."),
        ("Peta Arah Rambat Panas (Kompas)", 
         "• Mengukur penyebaran panas ke 8 arah (Utara, Timur, Selatan, Barat, dll).\n• Panah menunjukkan arah aliran panas utama pada telapak kaki."),
        ("Tabel Angka Lengkap (File CSV / Excel)", 
         "• Mencatat suhu tertinggi, terendah, dan rata-rata tiap titik kaki.\n• Skala ukuran fisik (px/cm) tercatat rapi untuk rekam medis."),
        ("Paket Arsip 1-Klik (File ZIP & JSON)", 
         "• Semua gambar, grafik, dan tabel dibungkus dalam 1 file ZIP.\n• File catatan (.json) bisa dibuka kembali kapan saja tanpa kehilangan data.")
    ]

    grid_w = Inches(5.66)
    grid_h = Inches(2.15)
    gx1 = Inches(0.8)
    gx2 = gx1 + grid_w + Inches(0.4)
    gy1 = Inches(2.2)
    gy2 = gy1 + grid_h + Inches(0.25)

    coords = [(gx1, gy1), (gx2, gy1), (gx1, gy2), (gx2, gy2)]

    for idx, (otitle, odesc) in enumerate(out_data):
        bx, by = coords[idx]
        b = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, grid_w, grid_h)
        b.fill.solid()
        b.fill.fore_color.rgb = CARD_BG
        b.line.color.rgb = CARD_BORDER
        b.line.width = Pt(1)

        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)

        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = f"{idx+1}. {otitle}\n"
        r.font.name = "Arial"
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = TEXT_ACCENT

        p_body = tf.add_paragraph()
        p_body.space_before = Pt(4)
        r_body = p_body.add_run()
        r_body.text = odesc
        r_body.font.name = "Arial"
        r_body.font.size = Pt(13)
        r_body.font.color.rgb = TEXT_BODY

    # Save to working directory
    out_path_work = r"D:\~Ideas n Innovation\~~Taiwan\AU\Aditya Thermal\ThermalSight-App-1.1.0\ThermalSight_v180_Alur_Kerja.pptx"
    prs.save(out_path_work)
    print(f"Saved PPTX to {out_path_work}")

    # Also save to GitHub directory
    out_path_git = r"C:\Users\Pongo\OneDrive\Documents\GitHub\ThermalSight-App\ThermalSight_v180_Alur_Kerja.pptx"
    prs.save(out_path_git)
    print(f"Saved PPTX to {out_path_git}")

if __name__ == '__main__':
    create_presentation()
